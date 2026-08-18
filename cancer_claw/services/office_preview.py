"""从 OOXML Office 文件抽出可预览的文本 / 表格 / 幻灯片。

不引入 python-docx / openpyxl：docx、xlsx 用 zip+xml；pptx 复用已有 python-pptx。
旧版 .doc / .xls / .ppt 是 OLE 二进制，这里明确拒绝。
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

_DOCX_EXTS = {".docx", ".docm"}
_XLSX_EXTS = {".xlsx", ".xlsm"}
_PPTX_EXTS = {".pptx", ".pptm"}
_LEGACY_EXTS = {".doc", ".xls", ".ppt"}

_MAX_DOCX_BLOCKS = 240
_MAX_XLSX_SHEETS = 6
_MAX_XLSX_COLS = 24
_MAX_PPTX_SLIDES = 40


def office_format(path: Path) -> str | None:
    ext = path.suffix.lower()
    if ext in _DOCX_EXTS:
        return "docx"
    if ext in _XLSX_EXTS:
        return "xlsx"
    if ext in _PPTX_EXTS:
        return "pptx"
    if ext in _LEGACY_EXTS:
        return "legacy"
    return None


def preview_office(
    path: Path,
    *,
    max_chars: int = 200_000,
    max_rows: int = 80,
) -> dict[str, Any]:
    kind = office_format(path)
    if kind == "legacy":
        raise ValueError(
            "旧版 .doc / .xls / .ppt 是二进制格式，无法在线预览。"
            "请另存为 .docx / .xlsx / .pptx 后查看，或直接下载。"
        )
    if kind == "docx":
        return _preview_docx(path, max_chars=max_chars)
    if kind == "xlsx":
        return _preview_xlsx(path, max_rows=max_rows)
    if kind == "pptx":
        return _preview_pptx(path, max_chars=max_chars)
    raise ValueError(f"不支持预览该 Office 格式: {path.suffix}")


def office_snippet(path: Path, max_chars: int = 4000) -> tuple[str, bool]:
    payload = preview_office(path, max_chars=max_chars, max_rows=12)
    text = payload.get("text") or ""
    truncated = bool(payload.get("truncated")) or len(text) > max_chars
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True
    return text, truncated


def _local(tag: str) -> str:
    return tag.split("}")[-1] if "}" in tag else tag


def _preview_docx(path: Path, *, max_chars: int) -> dict[str, Any]:
    paragraphs: list[str] = []
    tables: list[list[list[str]]] = []
    truncated = False
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml")
    except KeyError as e:
        raise ValueError("不是有效的 .docx（缺少 word/document.xml）") from e
    except zipfile.BadZipFile as e:
        raise ValueError("不是有效的 .docx（损坏的 zip）") from e

    root = ET.fromstring(xml)
    body = next((el for el in root.iter() if _local(el.tag) == "body"), root)
    used = 0
    blocks = 0
    for child in list(body):
        name = _local(child.tag)
        if name == "sectPr":
            continue
        if name == "tbl":
            table = _docx_table(child)
            if table:
                tables.append(table)
                used += sum(len(c) for row in table for c in row)
                blocks += 1
        else:
            line = _docx_para(child)
            if line:
                paragraphs.append(line)
                used += len(line)
                blocks += 1
        if used >= max_chars or blocks >= _MAX_DOCX_BLOCKS:
            truncated = True
            break

    text_parts = list(paragraphs)
    for table in tables:
        text_parts.append(_table_as_text(table))
    text = "\n\n".join(text_parts)
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True

    return {
        "kind": "docx",
        "path": path.name,
        "size": path.stat().st_size,
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "paragraphs": paragraphs,
        "tables": tables,
        "text": text,
        "truncated": truncated,
    }


def _docx_para(el: ET.Element) -> str:
    parts: list[str] = []
    for node in el.iter():
        if _local(node.tag) == "t" and node.text:
            parts.append(node.text)
        elif _local(node.tag) == "tab":
            parts.append("\t")
        elif _local(node.tag) == "br":
            parts.append("\n")
    return "".join(parts).strip()


def _docx_table(tbl: ET.Element) -> list[list[str]]:
    rows: list[list[str]] = []
    for tr in tbl:
        if _local(tr.tag) != "tr":
            continue
        cells: list[str] = []
        for tc in tr:
            if _local(tc.tag) != "tc":
                continue
            cells.append(_docx_para(tc))
        if cells:
            rows.append(cells)
    return rows


def _table_as_text(table: list[list[str]]) -> str:
    return "\n".join(" | ".join(row) for row in table)


def _preview_xlsx(path: Path, *, max_rows: int) -> dict[str, Any]:
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        raise ValueError("不是有效的 .xlsx（损坏的 zip）") from e

    try:
        shared = _xlsx_shared_strings(zf)
        sheets_meta = _xlsx_sheet_targets(zf)
        sheets: list[dict[str, Any]] = []
        any_truncated = False
        for name, target in sheets_meta[:_MAX_XLSX_SHEETS]:
            grid, truncated = _xlsx_sheet_grid(
                zf, target, shared, max_rows=max_rows, max_cols=_MAX_XLSX_COLS
            )
            any_truncated = any_truncated or truncated
            columns, rows = _grid_to_columns_rows(grid)
            sheets.append(
                {
                    "name": name,
                    "columns": columns,
                    "rows": rows,
                    "truncated": truncated,
                }
            )
        if not sheets:
            raise ValueError("工作簿里没有可读取的工作表")
        text_parts: list[str] = []
        for sheet in sheets:
            text_parts.append(f"## {sheet['name']}")
            header = sheet["columns"]
            if header:
                text_parts.append(" | ".join(header))
            for row in sheet["rows"][:12]:
                text_parts.append(" | ".join(row))
        return {
            "kind": "xlsx",
            "path": path.name,
            "size": path.stat().st_size,
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "sheets": sheets,
            "text": "\n".join(text_parts),
            "truncated": any_truncated or len(sheets_meta) > _MAX_XLSX_SHEETS,
        }
    finally:
        zf.close()


def _xlsx_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    try:
        raw = zf.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(raw)
    out: list[str] = []
    for si in root:
        if _local(si.tag) != "si":
            continue
        texts = [node.text or "" for node in si.iter() if _local(node.tag) == "t"]
        out.append("".join(texts))
    return out


def _xlsx_sheet_targets(zf: zipfile.ZipFile) -> list[tuple[str, str]]:
    try:
        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    except KeyError:
        names = [
            n for n in zf.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")
        ]
        return [(f"Sheet{i + 1}", n) for i, n in enumerate(sorted(names))]

    rid_to_target: dict[str, str] = {}
    for rel in rels:
        rid = rel.attrib.get("Id") or rel.attrib.get("id") or ""
        target = (rel.attrib.get("Target") or "").replace("\\", "/")
        if rid and target:
            if not target.startswith("xl/"):
                target = "xl/" + target.lstrip("/")
            rid_to_target[rid] = target

    out: list[tuple[str, str]] = []
    ns_r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for sheet in wb.iter():
        if _local(sheet.tag) != "sheet":
            continue
        name = sheet.attrib.get("name") or f"Sheet{len(out) + 1}"
        rid = sheet.attrib.get(f"{ns_r}id") or sheet.attrib.get("id") or ""
        target = rid_to_target.get(rid)
        if target:
            out.append((name, target))
    return out


def _col_row(ref: str) -> tuple[int, int]:
    i = 0
    col = 0
    while i < len(ref) and ref[i].isalpha():
        col = col * 26 + (ord(ref[i].upper()) - 64)
        i += 1
    row = int(ref[i:]) if i < len(ref) and ref[i:].isdigit() else 1
    return max(col - 1, 0), max(row - 1, 0)


def _xlsx_cell_value(cell: ET.Element, shared: list[str]) -> str:
    cell_type = cell.attrib.get("t", "")
    if cell_type == "s":
        v = next((n.text for n in cell if _local(n.tag) == "v"), None)
        try:
            return shared[int(v or "0")]
        except (ValueError, IndexError):
            return v or ""
    if cell_type == "inlineStr":
        texts = [n.text or "" for n in cell.iter() if _local(n.tag) == "t"]
        return "".join(texts)
    if cell_type == "b":
        v = next((n.text for n in cell if _local(n.tag) == "v"), "0")
        return "TRUE" if v == "1" else "FALSE"
    v = next((n.text for n in cell if _local(n.tag) == "v"), None)
    if v is not None:
        return v
    texts = [n.text or "" for n in cell.iter() if _local(n.tag) == "t"]
    return "".join(texts)


def _xlsx_sheet_grid(
    zf: zipfile.ZipFile,
    target: str,
    shared: list[str],
    *,
    max_rows: int,
    max_cols: int,
) -> tuple[list[list[str]], bool]:
    try:
        raw = zf.read(target)
    except KeyError:
        return [], False
    root = ET.fromstring(raw)
    grid: dict[tuple[int, int], str] = {}
    max_r = -1
    max_c = -1
    truncated = False
    for cell in root.iter():
        if _local(cell.tag) != "c":
            continue
        ref = cell.attrib.get("r") or ""
        if not ref:
            continue
        c, r = _col_row(ref)
        if r >= max_rows or c >= max_cols:
            truncated = True
            continue
        value = _xlsx_cell_value(cell, shared)
        if value == "":
            continue
        grid[(r, c)] = value
        if r > max_r:
            max_r = r
        if c > max_c:
            max_c = c
    if max_r < 0:
        return [], truncated
    rows: list[list[str]] = []
    for r in range(max_r + 1):
        rows.append([grid.get((r, c), "") for c in range(max_c + 1)])
    return rows, truncated


def _grid_to_columns_rows(grid: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    if not grid:
        return [], []
    header = grid[0]
    columns = [cell or f"列{i + 1}" for i, cell in enumerate(header)]
    return columns, grid[1:]


def _preview_pptx(path: Path, *, max_chars: int) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ValueError("缺少 python-pptx，无法预览 PPT") from e

    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    used = 0
    truncated = False
    for index, slide in enumerate(prs.slides, 1):
        if index > _MAX_PPTX_SLIDES:
            truncated = True
            break
        chunks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    chunks.append(text)
            if getattr(shape, "has_table", False):
                table = shape.table
                lines = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                if lines:
                    chunks.append("\n".join(lines))
        title = chunks[0] if chunks else f"幻灯片 {index}"
        body = "\n".join(chunks[1:] if len(chunks) > 1 else chunks)
        slides.append({"index": index, "title": title, "body": body})
        used += len(title) + len(body)
        if used >= max_chars:
            truncated = True
            break

    text = "\n\n".join(
        f"## 幻灯片 {s['index']}  {s['title']}\n{s['body']}".strip() for s in slides
    )
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True
    return {
        "kind": "pptx",
        "path": path.name,
        "size": path.stat().st_size,
        "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "slides": slides,
        "text": text,
        "truncated": truncated,
    }
